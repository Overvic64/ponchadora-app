"""Lee los libros 'Horario de Clases por profesor' (formato F03PEAC03: una hoja por profesor, con la
cuadricula lunes-viernes) y los asigna a las personas de la base de datos."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from core.attendance_import import normalizar_nombre

DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
_HORAS = re.compile(r"(\d{1,2}):(\d{2})\D+(\d{1,2}):(\d{2})")
_GRUPO = re.compile(r"^[A-Za-z]*\d+-\d+\w*$")  # IA1-1, IMT10-1, AUT1-4mixto, 1-4mixto
_AVISO = re.compile(r"\b(?:Entrada|Salida)\s+\d{1,2}:\d{2}\s*hrs?\.?", re.IGNORECASE)
_TITULO = re.compile(r"\b(?:[A-Za-z]{1,4}\.)+")  # M.S.C., ING., LIC., D.R. ...


@dataclass
class HorarioProfesor:
    hoja: str
    nombre: str
    carrera: str
    periodo: str
    categoria: str  # tipo de profesor segun la firma de la hoja ("" si no la trae)
    bloques: list[dict] = field(default_factory=list)


def clave_nombre(nombre: str) -> str:
    return normalizar_nombre(_TITULO.sub(" ", nombre))


def _bloque(dia: int, inicio: str, fin: str, texto: str) -> dict:
    avisos = _AVISO.findall(texto)
    lineas = [" ".join(_AVISO.sub(" ", linea).split()) for linea in texto.splitlines()]
    lineas = [x for x in lineas if x]
    grupo = lineas[0] if len(lineas) > 1 and _GRUPO.match(lineas[0]) else None
    materia = " ".join(lineas[1:] if grupo else lineas)
    return {
        "dia": dia, "hora_inicio": inicio, "hora_fin": fin, "tipo": "clase" if grupo else "apoyo",
        "grupo": grupo, "materia": materia, "nota": " / ".join(avisos) or None,
    }


def _leer_hoja(ws, carrera: str) -> HorarioProfesor | None:
    nombre = categoria = periodo = ""
    cabecera = None
    for fila in ws.iter_rows(min_col=1, max_col=7):
        a, b = (str(fila[0].value or "").strip(), str(fila[1].value or "").strip())
        if a.startswith("Profesor"):
            nombre = b
        elif a.startswith("Cuatrimestre:"):
            periodo = a.split(":", 1)[1].strip()
        elif cabecera is None and b.upper() == "LUNES":
            cabecera = fila[0].row
        elif b.upper().startswith("PROFESOR"):  # firma: PROFESOR/A DE TIEMPO COMPLETO o DE ASIGNATURA
            categoria = "Profesor de Tiempo Completo" if "TIEMPO COMPLETO" in b.upper() else "Profesor de Asignatura"
    if not nombre or nombre.lower() == "nombre" or cabecera is None:
        return None  # hoja en blanco (plantilla)

    franjas = {}  # fila -> (inicio, fin)
    for fila in range(cabecera + 1, ws.max_row + 1):
        m = _HORAS.search(str(ws.cell(fila, 1).value or ""))
        if m:
            franjas[fila] = (f"{int(m[1]):02d}:{m[2]}", f"{int(m[3]):02d}:{m[4]}")
    combinadas = {(r.min_row, r.min_col): r.max_row for r in ws.merged_cells.ranges}
    horario = HorarioProfesor(ws.title, nombre, carrera, periodo, categoria)
    for fila, (inicio, _) in franjas.items():
        for col in range(2, 7):
            texto = str(ws.cell(fila, col).value or "").strip()
            if not texto:
                continue
            ultima = max(f for f in franjas if f <= combinadas.get((fila, col), fila))
            horario.bloques.append(_bloque(col - 1, inicio, franjas[ultima][1], texto))
    horario.bloques.sort(key=lambda b: (b["dia"], b["hora_inicio"]))
    return horario


def leer_horarios(ruta: str | Path, carrera: str) -> list[HorarioProfesor]:
    """Un horario por profesor. Si una persona tiene dos hojas (version corregida), gana la ultima.
    Las hojas sin clases ni actividades (ej. profesores sin carga) se omiten."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    por_persona: dict[str, HorarioProfesor] = {}
    for hoja in libro.worksheets:
        horario = _leer_hoja(hoja, carrera)
        if horario and horario.bloques:
            por_persona[clave_nombre(horario.nombre)] = horario
    return list(por_persona.values())


def emparejar(horarios: list[HorarioProfesor], maestros) -> tuple[list[tuple[HorarioProfesor, object]], list[HorarioProfesor]]:
    """Une cada horario con su persona: nombre igual sin importar orden/acentos/titulos; si no, el nombre de
    uno contenido en el del otro (ej. 'Machado Garrobo Alejandra' vs 'Jazmin Alejandra Machado Garrobo'),
    siempre que haya una sola candidata. Regresa (parejas, sin pareja)."""
    por_clave: dict[str, list] = {}
    for m in maestros:
        por_clave.setdefault(normalizar_nombre(m.nombre), []).append(m)
    parejas, sin_pareja = [], []
    for h in horarios:
        palabras = set(clave_nombre(h.nombre).split())
        candidatas = por_clave.get(" ".join(sorted(palabras))) or [
            m for m in maestros
            if len(palabras) >= 3 and len(set(normalizar_nombre(m.nombre).split())) >= 3
            and (set(normalizar_nombre(m.nombre).split()) <= palabras or palabras <= set(normalizar_nombre(m.nombre).split()))
        ]
        if len(candidatas) == 1:
            parejas.append((h, candidatas[0]))
        else:
            sin_pareja.append(h)
    return parejas, sin_pareja


def carrera_del_archivo(ruta: str | Path) -> str | None:
    """'Horario de Clases por profesor TI Sep-Dic 2026.xlsx' -> 'TI'."""
    m = re.search(r"profesor\s+(\S+)", Path(ruta).stem, re.IGNORECASE)
    return m[1].upper() if m else None


def importar(archivos: list[tuple[str | Path, str]], store) -> dict:
    """Lee cada (archivo, carrera) y guarda el horario de quien se reconozca en la base de datos. Se puede repetir:
    reemplaza el horario de esa persona/carrera/periodo en vez de duplicarlo. 'cambian_tipo' lista a quienes la hoja
    firma con un tipo distinto al registrado (solo se informa, no se modifica)."""
    asignados, sin_pareja, cambian_tipo = [], [], []
    for ruta, carrera in archivos:
        parejas, sin = emparejar(leer_horarios(ruta, carrera), store.maestros)
        sin_pareja += [f"{h.nombre} ({carrera})" for h in sin]
        for h, maestro in parejas:
            store.reemplazar_horarios(maestro.id, carrera, h.periodo, h.bloques)
            asignados.append(f"{maestro.nombre} ({carrera})")
            if h.categoria and maestro.tipo_personal and h.categoria != maestro.tipo_personal:
                cambian_tipo.append(f"{maestro.nombre}: base = {maestro.tipo_personal}, hoja = {h.categoria}")
    return {"asignados": asignados, "sin_pareja": sin_pareja, "cambian_tipo": sorted(set(cambian_tipo))}
